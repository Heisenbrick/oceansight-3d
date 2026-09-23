# ==============================================================================
# OceanSight 3D - SIH26067 Master Presentation Builder (Pixel-Perfect Polish)
# Team 404 Founders (Newton School of Technology x S-VYASA University)
# Fixes:
#   - Aspect ratio preservation on ALL images (logos, screenshots, diagrams)
#   - Framing & container cards for all images (no raw cutoffs)
#   - Elimination of all negative / overflowing slide coordinates
#   - Mathematical grid alignment: 0.60" margins, equal card gaps, aligned heights
#   - Fixed layering & z-order (timeline connector visible on top, no buried elements)
#   - Text embedded directly in shape text_frames for zero displacement
#   - Official SIH logo badge formatted with true 7.82 aspect ratio
# ==============================================================================

import os, sys, shutil
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ── Paths ─────────────────────────────────────────────────────────────────────
BASE_DIR     = r"C:\Users\ayush\.gemini\antigravity\scratch\oceansight-3d-mvp"
ASSETS_DIR   = os.path.join(BASE_DIR, "pipeline", "assets")
BRAIN_DIR    = r"C:\Users\ayush\.gemini\antigravity\brain\f7021cb9-7db4-4f07-a064-b5ca2192c170"
UPLOADED_DIR = os.path.join(BRAIN_DIR, ".user_uploaded")

IMG_SIH        = os.path.join(ASSETS_DIR, "sih_logo.png")
IMG_NASA       = os.path.join(ASSETS_DIR, "nasa_logo.png")
IMG_COPERNICUS = os.path.join(ASSETS_DIR, "copernicus_logo.png")
IMG_NOAA       = os.path.join(ASSETS_DIR, "noaa_logo.png")
IMG_PARAVIEW   = os.path.join(ASSETS_DIR, "paraview_logo.png")

IMG_PANOPLY_SCREEN = os.path.join(UPLOADED_DIR, "media_1789155874882.png")
IMG_MYOCEAN_SCREEN = os.path.join(UPLOADED_DIR, "media_1789155798457.png")
IMG_DEMO_SURFACE   = os.path.join(BRAIN_DIR, "demo_surface.jpg")
IMG_DEMO_THERMO    = os.path.join(BRAIN_DIR, "demo_thermocline.jpg")
IMG_ARCH_DIAGRAM   = os.path.join(BRAIN_DIR, "ocean_architecture_pipeline_1788867700886.jpg")

# ── Color Palette ─────────────────────────────────────────────────────────────
NAVY_DEEP = RGBColor(0x04, 0x0C, 0x1E)   # Deep ocean midnight canvas
NAVY_MID  = RGBColor(0x0A, 0x18, 0x2E)   # Card background
NAVY_CARD = RGBColor(0x0F, 0x22, 0x3D)   # Raised element fill
TEAL_DEEP = RGBColor(0x0A, 0x48, 0x5C)   # Oceanic teal accent
TEAL_BRT  = RGBColor(0x00, 0x8C, 0xA0)   # Vivid cyan-teal
CYAN      = RGBColor(0x38, 0xBD, 0xF8)   # Bright cyan highlight
CYAN_SOFT = RGBColor(0xBA, 0xE6, 0xFD)   # Light cyan text
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)   # Crisp white
OFFWHITE  = RGBColor(0xF1, 0xF5, 0xF9)   # Slate offwhite
MUTED     = RGBColor(0x94, 0xA3, 0xB8)   # Muted body slate
AMBER     = RGBColor(0xF5, 0x9E, 0x0B)   # Warning / Phase amber
CORAL     = RGBColor(0xF4, 0x3F, 0x5E)   # Alert / Problem coral
GREEN     = RGBColor(0x10, 0xB9, 0x81)   # Feasibility / Success green
PURPLE    = RGBColor(0xA7, 0x8B, 0xFA)   # Secondary accent

SW = 13.333   # 16:9 Widescreen width in inches
SH = 7.500    # 16:9 Widescreen height in inches

prs = Presentation()
prs.slide_width  = Inches(SW)
prs.slide_height = Inches(SH)
blank_layout = prs.slide_layouts[6]

# ── Helper Functions ──────────────────────────────────────────────────────────

def create_shape(slide, shape_type, x, y, w, h, fill=None, line_color=None, line_width=Pt(1), radius=None):
    """Creates a basic shape with explicit geometry and styling."""
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line_color:
        s.line.color.rgb = line_color
        s.line.width = line_width
    else:
        s.line.fill.background()
    if radius is not None and len(s.adjustments) > 0:
        s.adjustments[0] = radius
    return s

def add_card(slide, x, y, w, h, fill=NAVY_MID, border=CYAN, border_w=Pt(1), radius=0.08):
    """Rounded card container for grouping content cleanly."""
    return create_shape(slide, 5, x, y, w, h, fill=fill, line_color=border, line_width=border_w, radius=radius)

def add_fitted_picture(slide, img_path, box_x, box_y, box_w, box_h, pad_x=0.0, pad_y=0.0):
    """
    Fits image inside bounding box (box_x, box_y, box_w, box_h) maintaining its exact
    aspect ratio, centered horizontally and vertically. No stretching, zero distortion.
    """
    if not os.path.exists(img_path):
        return None
    try:
        with Image.open(img_path) as im:
            iw, ih = im.size
    except Exception:
        return None

    avail_w = box_w - (2.0 * pad_x)
    avail_h = box_h - (2.0 * pad_y)
    if avail_w <= 0 or avail_h <= 0 or iw <= 0 or ih <= 0:
        return None

    scale = min(avail_w / iw, avail_h / ih)
    tw = iw * scale
    th = ih * scale
    tx = box_x + pad_x + (avail_w - tw) / 2.0
    ty = box_y + pad_y + (avail_h - th) / 2.0

    return slide.shapes.add_picture(img_path, Inches(tx), Inches(ty), width=Inches(tw), height=Inches(th))

def add_text_box(slide, text, x, y, w, h, size=14, bold=False, color=WHITE,
                 align=PP_ALIGN.LEFT, italic=False, wrap=True, line_spacing=Pt(4)):
    """Simple text box with single paragraph."""
    txb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    txb.word_wrap = wrap
    tf = txb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return txb

def set_shape_text(shape, text, size=11, bold=False, color=WHITE, align=PP_ALIGN.CENTER):
    """Sets text directly on a shape's built-in text frame with vertical centering."""
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color

def add_sih_header_badge(slide):
    """
    Renders the official SIH logo and problem ID at top-right.
    Preserves exact 7.82 aspect ratio of SIH logo without any stretching.
    """
    badge_w = 3.65
    badge_h = 0.62
    badge_x = SW - badge_w - 0.40   # 0.40" right margin
    badge_y = 0.165

    # White pill container card
    bg = create_shape(slide, 5, badge_x, badge_y, badge_w, badge_h,
                      fill=WHITE, line_color=CYAN, line_width=Pt(1.5), radius=0.25)

    # Insert official SIH logo maintaining true proportions
    if os.path.exists(IMG_SIH):
        add_fitted_picture(slide, IMG_SIH, badge_x + 0.10, badge_y + 0.05, 2.45, badge_h - 0.10)

    # Sub-badge with PS ID
    tag = create_shape(slide, 5, badge_x + 2.65, badge_y + 0.08, 0.90, badge_h - 0.16,
                       fill=NAVY_DEEP, line_color=None, radius=0.30)
    tf = tag.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "SIH 26067"
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = CYAN

def add_header_banner(slide, title, subtitle=None):
    """Uniform modern header spanning full width with breadcrumbs and logo."""
    # Top banner bar
    create_shape(slide, 1, 0, 0, SW, 0.95, fill=NAVY_MID, line_color=None)
    # Accent cyan hairline under header
    create_shape(slide, 1, 0, 0.95, SW, 0.03, fill=TEAL_BRT, line_color=None)

    # Title
    add_text_box(slide, title, 0.60, 0.10, 8.50, 0.48, size=21, bold=True, color=WHITE)
    # Subtitle / Category
    if subtitle:
        add_text_box(slide, subtitle, 0.60, 0.56, 8.50, 0.32, size=10.5, color=CYAN)

    add_sih_header_badge(slide)

def add_safe_marine_decor(slide):
    """Subtle oceanic accent bubbles strictly clamped inside slide margins."""
    dots = [(0.30, 0.30), (1.10, 0.20), (12.20, 0.20), (0.35, 7.15), (12.85, 7.15)]
    for dx, dy in dots:
        create_shape(slide, 9, dx, dy, 0.12, 0.08, fill=TEAL_BRT, line_color=None)


# ==============================================================================
# SLIDE 1: TITLE PAGE (Bold Ocean Hero with Pixel-Perfect Framed Visuals)
# ==============================================================================
print("[1/7] Building Slide 1: Title Page...")
s1 = prs.slides.add_slide(blank_layout)
create_shape(s1, 1, 0, 0, SW, SH, fill=NAVY_DEEP, line_color=None)
add_safe_marine_decor(s1)

# Top brand color line
create_shape(s1, 1, 0, 0, SW, 0.12, fill=TEAL_BRT, line_color=None)

# ── Left Column: Typography & Team Credentials ────────────────────────────────
LEFT_X = 0.60
LEFT_W = 6.20

# Hackathon Pill
pill = create_shape(s1, 5, LEFT_X, 0.60, 4.40, 0.38, fill=TEAL_DEEP, line_color=CYAN, line_width=Pt(1), radius=0.5)
set_shape_text(pill, "SMART INDIA HACKATHON 2024  ·  PS ID: SIH26067", size=9.5, bold=True, color=WHITE)

# Big Title
add_text_box(s1, "OCEANSIGHT 3D", LEFT_X, 1.15, LEFT_W, 0.85, size=40, bold=True, color=CYAN)

# Subtitle
add_text_box(s1, "Interactive 3D Visualization Platform for\nNumerical Ocean Models & In-Situ Observations",
             LEFT_X, 2.05, LEFT_W, 0.85, size=15, bold=False, color=WHITE)

# Team Card
t_card = add_card(s1, LEFT_X, 3.10, LEFT_W, 2.60, fill=NAVY_MID, border=TEAL_BRT, border_w=Pt(1.2), radius=0.08)

add_text_box(s1, "Team 404 Founders", LEFT_X + 0.25, 3.25, LEFT_W - 0.50, 0.45, size=21, bold=True, color=WHITE)
add_text_box(s1, "Newton School of Technology  ×  S-VYASA University",
             LEFT_X + 0.25, 3.75, LEFT_W - 0.50, 0.35, size=12.5, color=CYAN_SOFT)
add_text_box(s1, "Ministry of Earth Sciences (MoES)  |  INCOIS",
             LEFT_X + 0.25, 4.15, LEFT_W - 0.50, 0.35, size=11.5, bold=True, color=CYAN)

# Hairline divider
create_shape(s1, 1, LEFT_X + 0.25, 4.60, LEFT_W - 0.50, 0.02, fill=TEAL_BRT, line_color=None)

add_text_box(s1, "Theme: Research, Innovation & Operational Maritime Safety (Disaster Mgmt)",
             LEFT_X + 0.25, 4.75, LEFT_W - 0.50, 0.40, size=9.5, italic=True, color=MUTED)

add_text_box(s1, "Category: Software  |  Domain: Oceanography & Cloud-Native 3D GIS",
             LEFT_X + 0.25, 5.15, LEFT_W - 0.50, 0.35, size=9.5, color=MUTED)

# 3 Status Badges
b_w = (LEFT_W - 0.30) / 3.0
badges = ["✓ Live WebGL 3D", "✓ NetCDF → Zarr Pipeline", "✓ Zero-Install Web"]
for i, b_text in enumerate(badges):
    bx = LEFT_X + i * (b_w + 0.15)
    b_shp = create_shape(s1, 5, bx, 5.95, b_w, 0.42, fill=NAVY_CARD, line_color=CYAN, line_width=Pt(1), radius=0.4)
    set_shape_text(b_shp, b_text, size=9, bold=True, color=CYAN)

# ── Right Column: Framed 3D Engine & Thermocline Showcase ────────────────────
RIGHT_X = 7.15
RIGHT_W = 5.58

# Top Card: 3D Surface View (16:9 proportion, framed)
c1_h = 3.30
c1 = add_card(s1, RIGHT_X, 0.60, RIGHT_W, c1_h, fill=NAVY_MID, border=CYAN, border_w=Pt(1.5), radius=0.06)
if os.path.exists(IMG_DEMO_SURFACE):
    add_fitted_picture(s1, IMG_DEMO_SURFACE, RIGHT_X, 0.60, RIGHT_W, c1_h - 0.45, pad_x=0.12, pad_y=0.10)
# Caption strip at bottom of card 1
cap1 = create_shape(s1, 1, RIGHT_X + 0.02, 0.60 + c1_h - 0.40, RIGHT_W - 0.04, 0.38, fill=NAVY_DEEP, line_color=None)
set_shape_text(cap1, "🌊 Volumetric Thermocline & Gerstner Waves (Three.js WebGL Engine)", size=9.5, bold=True, color=CYAN)

# Bottom Card: Thermocline Depth Profile
c2_h = 2.45
c2_y = 4.10
c2 = add_card(s1, RIGHT_X, c2_y, RIGHT_W, c2_h, fill=NAVY_MID, border=TEAL_BRT, border_w=Pt(1.2), radius=0.06)
if os.path.exists(IMG_DEMO_THERMO):
    add_fitted_picture(s1, IMG_DEMO_THERMO, RIGHT_X, c2_y, RIGHT_W, c2_h - 0.40, pad_x=0.12, pad_y=0.08)
cap2 = create_shape(s1, 1, RIGHT_X + 0.02, c2_y + c2_h - 0.38, RIGHT_W - 0.04, 0.36, fill=NAVY_DEEP, line_color=None)
set_shape_text(cap2, "📍 Subsurface Acoustic Shadow Slices (0m → 500m Continuous Depth)", size=9.5, bold=True, color=CYAN_SOFT)

# Top Right SIH Logo
add_sih_header_badge(s1)


# ==============================================================================
# SLIDE 2: THE CORE PROBLEM & OPERATIONAL PAIN POINTS
# ==============================================================================
print("[2/7] Building Slide 2: Core Problem & Pain Points...")
s2 = prs.slides.add_slide(blank_layout)
create_shape(s2, 1, 0, 0, SW, SH, fill=NAVY_DEEP, line_color=None)
add_safe_marine_decor(s2)
add_header_banner(s2, "THE CORE PROBLEM & OPERATIONAL PAIN POINTS",
                  "Problem Statement ID: SIH26067 | Ministry of Earth Sciences / INCOIS")

# Problem Statement Quote Box (Sleek cyan highlight)
q_box = create_shape(s2, 5, 0.60, 1.15, 12.13, 0.70, fill=NAVY_MID, line_color=TEAL_BRT, line_width=Pt(1.2), radius=0.15)
set_shape_text(q_box,
               '"Develop a web-based interactive 3D visualization platform that integrates numerical ocean model outputs and in-situ observations."',
               size=12.5, bold=True, color=CYAN)

# 4 Pain-Point Cards
col_w = 2.88
col_gap = 0.20
pain_cards = [
    ("📊 Data Overload vs. Insight", CORAL, [
        "• INCOIS generates massive multi-dimensional NetCDF/GRIB outputs daily.",
        "• Interpreting 4D ocean physics requires doctorate-level expertise.",
        "• Raw numbers conceal critical operational gradients like thermoclines and shear currents."
    ]),
    ("🚢 Operational Users at Sea", AMBER, [
        "• Coast Guard, Naval Submarines, Cargo Captains, and ROVs depend on ocean state daily.",
        "• Submarines require exact acoustic shadow depths for stealth navigation.",
        "• Subsea maintenance robots struggle with sudden deep current shear."
    ]),
    ("⏳ Manual Workflow Bottleneck", TEAL_BRT, [
        "• Scientists must open complex legacy desktop tools (Panoply, Ferret, Ncview).",
        "• Slices and graphs are manually compiled into static PDF reports.",
        "• Reports are emailed to captains, who then must call scientists for clarifications."
    ]),
    ("⚠️ Delayed Life-or-Death Decisions", CORAL, [
        "• Operators receive static 2D snapshots hours after models are run.",
        "• No real-time interactive exploration when conditions suddenly shift at sea.",
        "• Critical gap between scientific numerical models and fast operational decisions."
    ])
]

for i, (title, accent, bullets) in enumerate(pain_cards):
    cx = 0.60 + i * (col_w + col_gap)
    # Card container
    card = add_card(s2, cx, 2.05, col_w, 4.80, fill=NAVY_MID, border=accent, border_w=Pt(1.2), radius=0.08)

    # Title header strip
    h_strip = create_shape(s2, 5, cx + 0.08, 2.15, col_w - 0.16, 0.45, fill=NAVY_CARD, line_color=None, radius=0.2)
    set_shape_text(h_strip, title, size=11, bold=True, color=WHITE)

    # Bullet content
    body = "\n\n".join(bullets)
    add_text_box(s2, body, cx + 0.15, 2.75, col_w - 0.30, 3.90, size=9.8, color=OFFWHITE, wrap=True)


# ==============================================================================
# SLIDE 3: CURRENT LANDSCAPE: EXISTING TOOLS & PLATFORMS
# ==============================================================================
print("[3/7] Building Slide 3: Current Landscape & Tools...")
s3 = prs.slides.add_slide(blank_layout)
create_shape(s3, 1, 0, 0, SW, SH, fill=NAVY_DEEP, line_color=None)
add_safe_marine_decor(s3)
add_header_banner(s3, "CURRENT LANDSCAPE: EXISTING TOOLS & PLATFORMS",
                  "Reviewing world-wide tools currently used for NetCDF ocean visualization")

tool_data = [
    ("NASA Panoply", "USA · NASA GSFC", IMG_NASA, [
        "• The world's most widely used desktop viewer for NetCDF, HDF, and GRIB datasets.",
        "• Produces static 2D slice contour and vector maps.",
        "• Standard tool in atmospheric and academic oceanographic institutions."
    ]),
    ("Copernicus MyOcean Pro", "European Union · Mercator", IMG_COPERNICUS, [
        "• Flagship web-based GIS portal for EU marine data visualization.",
        "• Provides interactive 2D map layers, depth sliders, and point time-series.",
        "• Serves European global ocean forecast and reanalysis products."
    ]),
    ("NOAA ERDDAP / Ncview", "USA · NOAA CoastWatch", IMG_NOAA, [
        "• High-speed data server and lightweight NetCDF visual inspection utility.",
        "• Allows tabular filtering and fast slicing of gridded ocean data.",
        "• Widely adopted by research laboratories and data scientists."
    ]),
    ("ParaView / VisIt", "Global · Kitware HPC", IMG_PARAVIEW, [
        "• High-end desktop scientific 3D visualization and volume rendering suite.",
        "• Supports complex isosurfaces, streamline tracing, and volumetric shaders.",
        "• Demands dedicated workstation GPUs and extensive specialized training."
    ])
]

for i, (name, origin, logo_path, bullets) in enumerate(tool_data):
    cx = 0.60 + i * (col_w + col_gap)
    # Master card
    add_card(s3, cx, 1.20, col_w, 5.65, fill=NAVY_MID, border=TEAL_BRT, border_w=Pt(1.2), radius=0.08)

    # Logo Well: dedicated white container with rounded corners
    logo_well = create_shape(s3, 5, cx + 0.20, 1.35, col_w - 0.40, 1.15, fill=WHITE, line_color=None, radius=0.15)
    # Fitted picture inside logo well — NO distortion, maintains exact aspect ratio!
    if os.path.exists(logo_path):
        add_fitted_picture(s3, logo_path, cx + 0.20, 1.35, col_w - 0.40, 1.15, pad_x=0.15, pad_y=0.10)

    # Tool Name
    add_text_box(s3, name, cx + 0.10, 2.65, col_w - 0.20, 0.40, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Origin Tag Pill
    tag = create_shape(s3, 5, cx + 0.25, 3.10, col_w - 0.50, 0.32, fill=TEAL_DEEP, line_color=CYAN, line_width=Pt(1), radius=0.5)
    set_shape_text(tag, origin, size=8.5, bold=True, color=WHITE)

    # Bullet content
    body = "\n\n".join(bullets)
    add_text_box(s3, body, cx + 0.15, 3.60, col_w - 0.30, 3.10, size=9.8, color=OFFWHITE, wrap=True)


# ==============================================================================
# SLIDE 4: LIMITATIONS & BOTTLENECKS OF CURRENT TOOLS
# ==============================================================================
print("[4/7] Building Slide 4: Limitations of Current Tools...")
s4 = prs.slides.add_slide(blank_layout)
create_shape(s4, 1, 0, 0, SW, SH, fill=NAVY_DEEP, line_color=None)
add_safe_marine_decor(s4)
add_header_banner(s4, "LIMITATIONS & BOTTLENECKS OF CURRENT TOOLS",
                  "Why existing solutions fail in fast-paced operational maritime scenarios")

pane_w = 5.95
pane_gap = 0.23

# ── Left Column: NASA Panoply ─────────────────────────────────────────────────
p_x = 0.60
add_card(s4, p_x, 1.15, pane_w, 5.75, fill=NAVY_MID, border=CORAL, border_w=Pt(1.2), radius=0.08)

# Header
h1 = create_shape(s4, 5, p_x + 0.12, 1.25, pane_w - 0.24, 0.42, fill=NAVY_CARD, line_color=CORAL, line_width=Pt(1), radius=0.2)
set_shape_text(h1, "NASA PANOPLY — Desktop Java Bottleneck", size=11.5, bold=True, color=WHITE)

# Framed Screenshot
frame1 = create_shape(s4, 1, p_x + 0.15, 1.75, pane_w - 0.30, 2.30, fill=NAVY_DEEP, line_color=TEAL_BRT, line_width=Pt(1))
if os.path.exists(IMG_PANOPLY_SCREEN):
    add_fitted_picture(s4, IMG_PANOPLY_SCREEN, p_x + 0.15, 1.75, pane_w - 0.30, 2.30, pad_x=0.08, pad_y=0.08)

# Drawbacks
panoply_drawbacks = [
    "❌ Desktop-Bound (Java): Requires downloading 15–20 GB raw files locally and installing JRE. Impossible on mobile or ship VSAT links.",
    "❌ Flat 2D Slices Only: View only one horizontal or vertical cross-section at a time. Zero true 3D volumetric rendering.",
    "❌ Zero In-Situ Sensor Overlay: Cannot overlay real-time ARGO floats or buoy observations alongside model fields.",
    "❌ High Cognitive Barrier: Designed strictly for atmospheric research PhDs; completely unusable for an operational Coast Guard captain."
]
p_text = "\n\n".join(panoply_drawbacks)
add_text_box(s4, p_text, p_x + 0.20, 4.15, pane_w - 0.40, 2.65, size=9.5, color=OFFWHITE, wrap=True)

# ── Right Column: Copernicus MyOcean Pro ──────────────────────────────────────
m_x = 0.60 + pane_w + pane_gap
add_card(s4, m_x, 1.15, pane_w, 5.75, fill=NAVY_MID, border=AMBER, border_w=Pt(1.2), radius=0.08)

# Header
h2 = create_shape(s4, 5, m_x + 0.12, 1.25, pane_w - 0.24, 0.42, fill=NAVY_CARD, line_color=AMBER, line_width=Pt(1), radius=0.2)
set_shape_text(h2, "COPERNICUS MyOcean PRO — 2.5D Surface Limitation", size=11.5, bold=True, color=WHITE)

# Framed Screenshot
frame2 = create_shape(s4, 1, m_x + 0.15, 1.75, pane_w - 0.30, 2.30, fill=NAVY_DEEP, line_color=TEAL_BRT, line_width=Pt(1))
if os.path.exists(IMG_MYOCEAN_SCREEN):
    add_fitted_picture(s4, IMG_MYOCEAN_SCREEN, m_x + 0.15, 1.75, pane_w - 0.30, 2.30, pad_x=0.08, pad_y=0.08)

# Drawbacks
myocean_drawbacks = [
    "❌ 2.5D Surface Globe, Not Volumetric: Paints color rasters onto the surface of a globe. Cannot slice depth volumes (0m → 500m).",
    "❌ No Acoustic / Thermocline Defense Focus: Fails to highlight acoustic shadow zones for submarine sonar or current shear for ROVs.",
    "❌ Eurocentric / Disconnected from India: Tailored to European regions; lacks native support for INCOIS high-res ROMS grids.",
    "❌ Complex Layering Hierarchy: Overwhelming GIS menus with high latency when toggling multiple variables."
]
m_text = "\n\n".join(myocean_drawbacks)
add_text_box(s4, m_text, m_x + 0.20, 4.15, pane_w - 0.40, 2.65, size=9.5, color=OFFWHITE, wrap=True)


# ==============================================================================
# SLIDE 5: OUR SOLUTION: PURPOSE-DRIVEN HYBRID ARCHITECTURE
# ==============================================================================
print("[5/7] Building Slide 5: Purpose-Driven Architecture...")
s5 = prs.slides.add_slide(blank_layout)
create_shape(s5, 1, 0, 0, SW, SH, fill=NAVY_DEEP, line_color=None)
add_safe_marine_decor(s5)
add_header_banner(s5, "OUR SOLUTION: PURPOSE-DRIVEN HYBRID ARCHITECTURE",
                  "Solving cognitive overload through intelligent data layering and hybrid 2D-to-3D exploration")

# Philosophy Quote
phil = create_shape(s5, 5, 0.60, 1.15, 12.13, 0.65, fill=NAVY_MID, line_color=CYAN, line_width=Pt(1.2), radius=0.2)
set_shape_text(phil,
               '"The main problem is not just 2D vs. 3D — it is how vast multi-parameter ocean data is layered, ordered, and presented so anyone can understand it."',
               size=11.5, bold=True, color=CYAN)

col_w3 = 3.91
col_gap3 = 0.20

pillars = [
    ("1. Cognitive Layering Hierarchy", TEAL_BRT, [
        "• Structured Operational Depths: Data is organized into clean intuitive planes:",
        "  - Surface Layer: SST heatmap + Gerstner dynamic waves.",
        "  - Subsurface Volume: Continuous thermocline gradient & isotherms.",
        "  - Dynamics: 3D particle current streamlines with velocity colors.",
        "  - Bathymetry: Procedural seabed contours & continental shelf.",
        "  - Ground Truth: Live ARGO profiling floats & sensor telemetry."
    ]),
    ("2. Hybrid 2D-to-3D Workflow", CYAN, [
        "• Macro 2D GIS Map (Surveillance / Overview):",
        "  Familiar, lightweight 2D map for regional ocean surveillance without heavy graphics rendering overhead.",
        "• On-Demand 3D Volumetric Slicing (Deep Dive):",
        "  User box-selects any coordinates on the 2D map to instantly instantiate a deep 3D ocean volume (0m → 500m → seabed) with live probe telemetry."
    ]),
    ("3. Zero-Friction Web Accessibility", GREEN, [
        "• 100% Browser-Native WebGL: Zero installs, zero downloads, and NO mandatory login barrier.",
        "• Mobile & Low-Bandwidth Ready: Fully responsive on tactical tablets and ship bridge screens over low-speed satellite/VSAT.",
        "• Unified Platform: Scales seamlessly from university research to frontline Coast Guard missions."
    ])
]

for i, (title, accent, pts) in enumerate(pillars):
    cx = 0.60 + i * (col_w3 + col_gap3)
    add_card(s5, cx, 1.95, col_w3, 4.90, fill=NAVY_MID, border=accent, border_w=Pt(1.2), radius=0.08)

    # Header
    hdr = create_shape(s5, 5, cx + 0.12, 2.05, col_w3 - 0.24, 0.45, fill=NAVY_CARD, line_color=accent, line_width=Pt(1), radius=0.2)
    set_shape_text(hdr, title, size=12, bold=True, color=WHITE)

    # Body
    body = "\n\n".join(pts)
    add_text_box(s5, body, cx + 0.18, 2.65, col_w3 - 0.36, 4.05, size=9.8, color=OFFWHITE, wrap=True)


# ==============================================================================
# SLIDE 6: END-TO-END TECHNOLOGY STACK
# ==============================================================================
print("[6/7] Building Slide 6: End-to-End Technology Stack...")
s6 = prs.slides.add_slide(blank_layout)
create_shape(s6, 1, 0, 0, SW, SH, fill=NAVY_DEEP, line_color=None)
add_safe_marine_decor(s6)
add_header_banner(s6, "END-TO-END TECHNOLOGY STACK",
                  "Modern, high-performance cloud-native tools driving OceanSight 3D")

tiers = [
    ("01", "Data Pipeline & Ingestion", CYAN, [
        "Python (Xarray, Dask)",
        "NetCDF4 Data Slicing",
        "Zarr Cloud Conversion",
        "INCOIS ROMS Grids",
        "ARGO Float In-Situ Feeds"
    ]),
    ("02", "Storage & Database", GREEN, [
        "PostgreSQL + PostGIS",
        "Geospatial Buoy Indexing",
        "Cloud Object Store (S3)",
        "Chunked Zarr Store",
        "Bathymetry DEM Cache"
    ]),
    ("03", "High-Speed API & Cache", AMBER, [
        "FastAPI (Async Microservice)",
        "Redis In-Memory Cache",
        "< 1ms Pre-warmed Slices",
        "Raw Binary Data Packing",
        "Cloudflare Edge CDN"
    ]),
    ("04", "Client 3D & 2D Frontend", CORAL, [
        "React.js + TailwindCSS",
        "Three.js / WebGL & WebGPU",
        "Custom GLSL Wave Shaders",
        "MapLibre GL (2D GIS)",
        "Mobile & Touch Responsive"
    ])
]

# 1. First add the 4 background cards
for i, (num, name, col, pts) in enumerate(tiers):
    cx = 0.60 + i * (col_w + col_gap)
    add_card(s6, cx, 1.20, col_w, 5.65, fill=NAVY_MID, border=col, border_w=Pt(1.2), radius=0.08)

# 2. Add the timeline connector line ON TOP of cards so it remains VISIBLE!
create_shape(s6, 1, 0.60 + 0.50, 3.25, 12.13 - 1.0, 0.04, fill=CYAN, line_color=None)

# 3. Add details inside each card
for i, (num, name, col, pts) in enumerate(tiers):
    cx = 0.60 + i * (col_w + col_gap)

    # Number Circle Badge
    nb = create_shape(s6, 9, cx + (col_w - 0.75)/2.0, 1.35, 0.75, 0.75, fill=col, line_color=WHITE, line_width=Pt(1.5))
    set_shape_text(nb, num, size=15, bold=True, color=NAVY_DEEP)

    # Tier Title
    add_text_box(s6, name, cx + 0.10, 2.25, col_w - 0.20, 0.55, size=12.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    # Timeline Node Dot on top of timeline line
    dot = create_shape(s6, 9, cx + (col_w - 0.32)/2.0, 3.11, 0.32, 0.32, fill=col, line_color=WHITE, line_width=Pt(2))

    # 5 Tech Stack Pills with embedded text
    for j, pt in enumerate(pts):
        tag_y = 3.65 + j * 0.60
        tag = create_shape(s6, 5, cx + 0.15, tag_y, col_w - 0.30, 0.48,
                           fill=NAVY_CARD, line_color=col if j == 0 else TEAL_DEEP, line_width=Pt(1), radius=0.4)
        set_shape_text(tag, pt, size=9.5, bold=(j == 0), color=WHITE)


# ==============================================================================
# SLIDE 7: END-TO-END DATA CONVERSION & STREAMING ARCHITECTURE
# ==============================================================================
print("[7/7] Building Slide 7: Streaming Architecture & Pipeline...")
s7 = prs.slides.add_slide(blank_layout)
create_shape(s7, 1, 0, 0, SW, SH, fill=NAVY_DEEP, line_color=None)
add_safe_marine_decor(s7)
add_header_banner(s7, "END-TO-END DATA CONVERSION & STREAMING PIPELINE",
                  "NetCDF Ingestion → Python Zarr Conversion → Redis Caching → 3D WebGL Rendering")

# ── Left Column: Architecture Diagram Framed Card ────────────────────────────
diag_w = 7.70
diag_h = 4.65
diag_x = 0.60
diag_y = 1.15

add_card(s7, diag_x, diag_y, diag_w, diag_h, fill=NAVY_MID, border=CYAN, border_w=Pt(1.5), radius=0.08)
if os.path.exists(IMG_ARCH_DIAGRAM):
    add_fitted_picture(s7, IMG_ARCH_DIAGRAM, diag_x, diag_y, diag_w, diag_h, pad_x=0.15, pad_y=0.15)

# ── Right Column: 4 Pipeline Step Cards ───────────────────────────────────────
info_x = 8.55
info_w = 4.18

arch_steps = [
    ("1. Convert NetCDF to Zarr", AMBER,
     "Modern cloud-native format. Stores chunks independently so the browser fetches only the exact depth slice requested instead of downloading the entire 15 GB file."),
    ("2. In-Memory Redis Caching", GREEN,
     "Common depths (0m, 50m, 100m) are pre-warmed in RAM cache. Subsequent slice requests drop to sub-millisecond response latency (< 1ms)."),
    ("3. Raw Binary Streaming (10x Smaller)", CYAN,
     "Replaces bloated JSON with compact Float32 binary byte buffers, reducing network payload by 90% over shipboard satellite / VSAT links."),
    ("4. Level of Detail (LOD) Meshing", CORAL,
     "Dynamic downsampling: High resolution near the camera probe; lower mesh density in distant ocean zones for steady 60 FPS performance.")
]

for i, (title, col, desc) in enumerate(arch_steps):
    iy = 1.15 + i * 1.18
    # Step card
    add_card(s7, info_x, iy, info_w, 1.10, fill=NAVY_MID, border=col, border_w=Pt(1.2), radius=0.10)
    # Left accent indicator line
    create_shape(s7, 1, info_x + 0.08, iy + 0.12, 0.04, 0.86, fill=col, line_color=None)

    add_text_box(s7, title, info_x + 0.22, iy + 0.08, info_w - 0.35, 0.32, size=11, bold=True, color=col)
    add_text_box(s7, desc, info_x + 0.22, iy + 0.40, info_w - 0.35, 0.65, size=8.8, color=OFFWHITE, wrap=True)

# ── Bottom Row: 5 Feature Badges ──────────────────────────────────────────────
features = [
    ("✓ Zero Login Required", GREEN),
    ("✓ Mobile & VSAT Ready", CYAN),
    ("✓ NetCDF → Zarr: 10x Faster", AMBER),
    ("✓ Sub-ms Depth Slices", PURPLE),
    ("✓ Indian Ocean Tailored", CORAL)
]

pill_w = 2.26
pill_gap = 0.20
pill_y = 6.00
pill_h = 0.75

for i, (feat, col) in enumerate(features):
    px = 0.60 + i * (pill_w + pill_gap)
    fb = create_shape(s7, 5, px, pill_y, pill_w, pill_h, fill=NAVY_MID, line_color=col, line_width=Pt(1.2), radius=0.4)
    set_shape_text(fb, feat, size=10, bold=True, color=col)


# ==============================================================================
# SAVE PRESENTATION
# ==============================================================================
out_polished_dload = r"C:\Users\ayush\Downloads\OceanSight_SIH26067_Final_Polished.pptx"
out_polished_desk  = r"C:\Users\ayush\Desktop\OceanSight_SIH26067_Final_Polished.pptx"
out_final_dload     = r"C:\Users\ayush\Downloads\OceanSight_SIH26067_Final.pptx"
out_final_desk      = r"C:\Users\ayush\Desktop\OceanSight_SIH26067_Final.pptx"

# Always save to Polished files (guaranteed never locked)
prs.save(out_polished_dload)
prs.save(out_polished_desk)
print(f"\n[OK] Saved Polished Edition to:")
print(f"  - {out_polished_dload}")
print(f"  - {out_polished_desk}")

# Attempt saving to original files if PowerPoint lock allows
try:
    prs.save(out_final_dload)
    prs.save(out_final_desk)
    print(f"[OK] Successfully overwritten original files as well:")
    print(f"  - {out_final_dload}")
    print(f"  - {out_final_desk}")
except PermissionError:
    print(f"[NOTE] Original file '{out_final_dload}' is currently open in PowerPoint.")
    print(f"       You can view '{out_polished_dload}' directly, or close PowerPoint to overwrite.")

print("\nPresentation generation complete! All 7 slides upgraded to pixel-perfection.")
